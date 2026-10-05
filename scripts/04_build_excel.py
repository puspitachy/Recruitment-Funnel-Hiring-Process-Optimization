"""
04_build_excel.py
Builds excel/recruitment_funnel_analysis.xlsx — the Excel stage of the project.
Every summary number is a live formula (COUNTIFS / SUMIFS / AVERAGEIFS / INDEX-MATCH)
over the data sheets, so the workbook recalculates if the data changes.
Run scripts/recalc (LibreOffice) afterwards to store calculated values.
"""
import sqlite3
import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = Path(__file__).resolve().parents[1]
OUT = BASE / "excel"; OUT.mkdir(exist_ok=True)
con = sqlite3.connect(BASE / "data" / "recruitment.db")

# ------------------------------------------------------------------ styles
F = "Arial"
NAVY, BLUE_HEX, LIGHT, YELLOW = "1F3A5F", "2A78D6", "EEF3FA", "FFF2CC"
h_font = Font(name=F, bold=True, color="FFFFFF", size=10)
h_fill = PatternFill("solid", fgColor=NAVY)
title_font = Font(name=F, bold=True, size=16, color=NAVY)
sub_font = Font(name=F, italic=True, size=10, color="52514E")
bold = Font(name=F, bold=True, size=10)
norm = Font(name=F, size=10)
input_fill = PatternFill("solid", fgColor=YELLOW)
kpi_fill = PatternFill("solid", fgColor=LIGHT)
thin = Side(style="thin", color="C9C8C2")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
PCT, PCT2, INT, MONEY, DEC = "0.0%", "0.00%", "#,##0", "$#,##0", "0.0"

def title(ws, text, subtitle):
    ws["A1"] = text; ws["A1"].font = title_font
    ws["A2"] = subtitle; ws["A2"].font = sub_font
    ws.sheet_view.showGridLines = False

def header(ws, row, col, labels, widths=None):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font, c.fill, c.border = h_font, h_fill, box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(col + i)].width = w

def cell(ws, ref, value, fmt=None, font=norm, fill=None):
    c = ws[ref]; c.value = value; c.font = font; c.border = box
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    return c

def value_labels():
    """Data labels showing only the value (no series / category names)."""
    d = DataLabelList()
    d.showVal, d.showSerName, d.showCatName, d.showLegendKey, d.showPercent = True, False, False, False, False
    return d

def note(ws, ref, text):
    ws[ref] = text; ws[ref].font = sub_font

wb = Workbook()

# ------------------------------------------------------------------ data sheets
apps = pd.read_sql("SELECT * FROM vw_application_funnel", con)
apps["outcome_month"] = apps.outcome_date.str[:7]
app_cols = ["application_id", "candidate_id", "requisition_id", "department_name", "job_title", "job_level",
            "location", "recruiter_name", "source", "gender", "applied_date", "applied_month", "outcome",
            "outcome_reason", "outcome_date", "outcome_month", "reached_phone_screen", "reached_assessment",
            "reached_onsite", "reached_offer", "is_hired", "days_resume_screen", "days_phone_screen",
            "days_assessment", "days_onsite", "days_offer", "time_to_hire_days", "days_in_process",
            "expected_salary", "offered_salary", "offer_status", "decline_reason", "offer_vs_expected_pct",
            "start_date", "employment_status", "exit_date", "performance_rating_6m"]
apps = apps[app_cols].sort_values("applied_date")
for c in ["applied_date", "outcome_date", "start_date", "exit_date"]:
    apps[c] = pd.to_datetime(apps[c])
reqs = pd.read_sql("SELECT * FROM vw_requisition_summary", con)
for c in ["opened_date", "closed_date"]:
    reqs[c] = pd.to_datetime(reqs[c])
spend = pd.read_sql("SELECT * FROM source_spend", con)

def write_table(ws, df, name, start_row=1, widths=None, formula_cols=None):
    cols = list(df.columns) + [f[0] for f in (formula_cols or [])]
    header(ws, start_row, 1, cols)
    n = len(df)
    for r, row in enumerate(df.itertuples(index=False), start_row + 1):
        for c, v in enumerate(row, 1):
            if pd.isna(v):
                continue
            cl = ws.cell(row=r, column=c, value=v.to_pydatetime() if isinstance(v, pd.Timestamp) else v)
            if isinstance(v, pd.Timestamp):
                cl.number_format = "yyyy-mm-dd"
        for k, (_, ftemplate, fmt) in enumerate(formula_cols or []):
            cl = ws.cell(row=r, column=len(df.columns) + 1 + k, value=ftemplate.format(r=r))
            if fmt: cl.number_format = fmt
    ref = f"A{start_row}:{get_column_letter(len(cols))}{start_row + n}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    for i, col in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(col, max(11, min(24, len(col) + 2)))
    return {col: get_column_letter(i) for i, col in enumerate(cols, 1)}, start_row + 1, start_row + n

ws_a = wb.active; ws_a.title = "Applications_Data"
L = {c: get_column_letter(i) for i, c in enumerate(app_cols, 1)}
retained = ("=IF(AND(${hired}{{r}}=1,${start}{{r}}<>\"\",${start}{{r}}<=DATE(2024,12,31)),"
            "IF(OR(${exit}{{r}}=\"\",${exit}{{r}}-${start}{{r}}>=365),1,0),\"\")").format(
            hired=L["is_hired"], start=L["start_date"], exit=L["exit_date"])
A, a0, a1 = write_table(ws_a, apps, "Applications", formula_cols=[("retained_12m", retained, "0")])
A_rng = lambda col: f"Applications_Data!${A[col]}${a0}:${A[col]}${a1}"

ws_r = wb.create_sheet("Requisitions_Data")
R, r0, r1 = write_table(ws_r, reqs, "Requisitions")
R_rng = lambda col: f"Requisitions_Data!${R[col]}${r0}:${R[col]}${r1}"

ws_s = wb.create_sheet("Source_Spend")
S, s0, s1 = write_table(ws_s, spend, "SourceSpend", widths={"spend_type": 30})
for r in range(s0, s1 + 1):
    ws_s[f"{S['amount']}{r}"].number_format = MONEY
S_rng = lambda col: f"Source_Spend!${S[col]}${s0}:${S[col]}${s1}"

DEPTS = sorted(apps.department_name.unique())
SOURCES = ["Employee Referral", "Recruitment Agency", "LinkedIn", "Company Website", "Campus Recruiting",
           "Job Boards", "Social Media", "Unknown"]
RECS = sorted(apps.recruiter_name.unique())

# ------------------------------------------------------------------ lists (for dropdowns)
ws_l = wb.create_sheet("Lists")
ws_l["A1"], ws_l["B1"] = "Department", "Source"
for i, v in enumerate(["All"] + DEPTS, 2): ws_l[f"A{i}"] = v
for i, v in enumerate(["All"] + SOURCES, 2): ws_l[f"B{i}"] = v
ws_l.sheet_state = "hidden"

# ------------------------------------------------------------------ Funnel (interactive)
ws = wb.create_sheet("Funnel", 0)
title(ws, "Hiring Funnel", "Pick a department and source in the yellow cells - every number below recalculates.")
for ref, lab in (("A4", "Department"), ("A5", "Source")):
    ws[ref] = lab; ws[ref].font = bold
for ref in ("B4", "B5"):
    cell(ws, ref, "All", fill=input_fill, font=bold)
dv1 = DataValidation(type="list", formula1=f"=Lists!$A$2:$A${len(DEPTS) + 2}", allow_blank=False)
dv2 = DataValidation(type="list", formula1=f"=Lists!$B$2:$B${len(SOURCES) + 2}", allow_blank=False)
ws.add_data_validation(dv1); ws.add_data_validation(dv2); dv1.add("B4"); dv2.add("B5")
crit = (f'{A_rng("department_name")},IF($B$4="All","*",$B$4),'
        f'{A_rng("source")},IF($B$5="All","*",$B$5)')
header(ws, 7, 1, ["Stage", "Candidates", "% of previous stage", "% of applicants", "Drop-off vs previous"],
       [24, 14, 14, 14, 14])
stages = [("Applied", f"=COUNTIFS({crit})"),
          ("Phone Screen", f"=SUMIFS({A_rng('reached_phone_screen')},{crit})"),
          ("Skills Assessment", f"=SUMIFS({A_rng('reached_assessment')},{crit})"),
          ("Onsite Interview", f"=SUMIFS({A_rng('reached_onsite')},{crit})"),
          ("Offer", f"=SUMIFS({A_rng('reached_offer')},{crit})"),
          ("Hired", f"=SUMIFS({A_rng('is_hired')},{crit})")]
for i, (stg, f) in enumerate(stages):
    r = 8 + i
    cell(ws, f"A{r}", stg, font=bold); cell(ws, f"B{r}", f, INT)
    cell(ws, f"C{r}", "" if i == 0 else f"=IFERROR(B{r}/B{r-1},0)", PCT)
    cell(ws, f"D{r}", f"=IFERROR(B{r}/$B$8,0)", PCT2)
    cell(ws, f"E{r}", "" if i == 0 else f"=B{r-1}-B{r}", INT)
ws.conditional_formatting.add("B8:B13", DataBarRule(start_type="min", end_type="max", color=BLUE_HEX))
cell(ws, "A15", "Applications per hire", font=bold); cell(ws, "B15", "=IFERROR(B8/B13,0)", "#,##0")
cell(ws, "A16", "Offer acceptance rate", font=bold)
# acceptance = accepted / (accepted + declined)
cell(ws, "B16", f'=IFERROR(COUNTIFS({crit},{A_rng("offer_status")},"Accepted")/'
                f'(COUNTIFS({crit},{A_rng("offer_status")},"Accepted")+COUNTIFS({crit},{A_rng("offer_status")},"Declined")),0)', PCT)
cell(ws, "A17", "Avg time to hire (days)", font=bold)
cell(ws, "B17", f'=IFERROR(AVERAGEIFS({A_rng("time_to_hire_days")},{crit},{A_rng("is_hired")},1),0)', DEC)
ch = BarChart(); ch.type = "bar"; ch.style = 10; ch.title = "Candidates reaching each stage"
ch.add_data(Reference(ws, min_col=2, min_row=7, max_row=13), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=8, max_row=13))
ch.x_axis.scaling.orientation = "maxMin"; ch.legend = None; ch.height, ch.width = 8, 16
ch.series[0].graphicalProperties.solidFill = BLUE_HEX
ch.dataLabels = value_labels()
ws.add_chart(ch, "G4")

# stage outcomes table
header(ws, 20, 1, ["Stage", "Entered", "Passed", "Rejected", "Withdrew", "Pass rate", "Withdrawal rate"])
stage_flags = [("Phone Screen", "reached_phone_screen", "reached_assessment"),
               ("Skills Assessment", "reached_assessment", "reached_onsite"),
               ("Onsite Interview", "reached_onsite", "reached_offer")]
for i, (stg, here, nxt) in enumerate(stage_flags):
    r = 21 + i
    cell(ws, f"A{r}", stg, font=bold)
    cell(ws, f"B{r}", f"=SUMIFS({A_rng(here)},{crit})", INT)
    cell(ws, f"C{r}", f"=SUMIFS({A_rng(nxt)},{crit})", INT)
    cell(ws, f"D{r}", f'=COUNTIFS({crit},{A_rng("outcome")},"Rejected",'
                      f'{A_rng(here)},1,{A_rng(nxt)},0)', INT)
    cell(ws, f"E{r}", f'=COUNTIFS({crit},{A_rng("outcome")},"Withdrew",{A_rng(here)},1,{A_rng(nxt)},0)', INT)
    cell(ws, f"F{r}", f"=IFERROR(C{r}/B{r},0)", PCT)
    cell(ws, f"G{r}", f"=IFERROR(E{r}/B{r},0)", PCT)
note(ws, "A25", "Rejected includes candidates closed out because the role was filled or cancelled.")

# ------------------------------------------------------------------ Stage_Speed
ws = wb.create_sheet("Stage_Speed", 1)
title(ws, "Where Is the Process Slow?", "Average days spent in each stage, hired candidates only (AVERAGEIFS).")
stage_days = [("Resume Screen", "days_resume_screen"), ("Phone Screen", "days_phone_screen"),
              ("Skills Assessment", "days_assessment"), ("Onsite Interview", "days_onsite"), ("Offer", "days_offer")]
header(ws, 4, 1, ["Department"] + [s for s, _ in stage_days] + ["Avg time to hire", "Hires"],
       [22, 13, 13, 13, 13, 13, 14, 9])
for i, d in enumerate(DEPTS + ["All departments"]):
    r = 5 + i
    cell(ws, f"A{r}", d, font=bold)
    dcrit = "" if d == "All departments" else f',{A_rng("department_name")},$A{r}'
    for j, (_, col) in enumerate(stage_days):
        cell(ws, f"{get_column_letter(2 + j)}{r}",
             f'=IFERROR(AVERAGEIFS({A_rng(col)},{A_rng("is_hired")},1{dcrit}),0)', DEC)
    cell(ws, f"G{r}", f'=IFERROR(AVERAGEIFS({A_rng("time_to_hire_days")},{A_rng("is_hired")},1{dcrit}),0)', DEC)
    cell(ws, f"H{r}", f'=COUNTIFS({A_rng("is_hired")},1{dcrit})', INT)
last = 5 + len(DEPTS) - 1
ws.conditional_formatting.add(f"B5:F{last}", ColorScaleRule(start_type="min", start_color="F7FBFF",
                                                            end_type="max", end_color="2A78D6"))
for c in "ABCDEFGH":
    ws[f"{c}{last + 1}"].font = bold
note(ws, f"A{last + 3}", "Darker = slower. The Onsite Interview column is the bottleneck, worst in Engineering and Data & Analytics.")
ch = BarChart(); ch.type = "col"; ch.title = "Average days per stage (all departments)"; ch.style = 10
ch.add_data(Reference(ws, min_col=2, max_col=6, min_row=last + 1), from_rows=True, titles_from_data=False)
ch.set_categories(Reference(ws, min_col=2, max_col=6, min_row=4))
ch.legend = None; ch.height, ch.width = 7.5, 16
ch.series[0].graphicalProperties.solidFill = BLUE_HEX
ch.dataLabels = value_labels()
ws.add_chart(ch, f"A{last + 5}")

# ------------------------------------------------------------------ Source_Analysis
ws = wb.create_sheet("Source_Analysis", 2)
title(ws, "Which Sources Are Worth It?", "Volume, conversion, cost and quality by candidate source.")
cols = ["Source", "Applications", "Resume pass %", "Hires", "App-to-hire %", "Apps per hire",
        "Avg time to hire", "Total spend", "Cost per hire", "12-month retention", "Avg performance (6m)"]
header(ws, 4, 1, cols, [22, 13, 12, 9, 12, 12, 12, 14, 13, 13, 13])
for i, s in enumerate(SOURCES):
    r = 5 + i
    sc = f'{A_rng("source")},$A{r}'
    cell(ws, f"A{r}", s, font=bold)
    cell(ws, f"B{r}", f"=COUNTIFS({sc})", INT)
    cell(ws, f"C{r}", f"=IFERROR(SUMIFS({A_rng('reached_phone_screen')},{sc})/B{r},0)", PCT)
    cell(ws, f"D{r}", f"=SUMIFS({A_rng('is_hired')},{sc})", INT)
    cell(ws, f"E{r}", f"=IFERROR(D{r}/B{r},0)", PCT2)
    cell(ws, f"F{r}", f"=IFERROR(B{r}/D{r},0)", INT)
    cell(ws, f"G{r}", f"=IFERROR(AVERAGEIFS({A_rng('time_to_hire_days')},{sc},{A_rng('is_hired')},1),0)", DEC)
    cell(ws, f"H{r}", f"=SUMIFS({S_rng('amount')},{S_rng('source')},$A{r})", MONEY)
    cell(ws, f"I{r}", f'=IF(H{r}=0,"n/a",IFERROR(H{r}/D{r},0))', MONEY)
    cell(ws, f"J{r}", f'=IF(D{r}<10,"too few",IFERROR(AVERAGEIFS(Applications_Data!${A["retained_12m"]}${a0}:${A["retained_12m"]}${a1},{sc}),0))', PCT)
    cell(ws, f"K{r}", f'=IF(D{r}<10,"too few",IFERROR(AVERAGEIFS({A_rng("performance_rating_6m")},{sc}),0))', "0.00")
lr = 5 + len(SOURCES)
cell(ws, f"A{lr}", "Total", font=bold)
for c, f, fmt in [("B", f"=SUM(B5:B{lr-1})", INT), ("D", f"=SUM(D5:D{lr-1})", INT),
                  ("E", f"=IFERROR(D{lr}/B{lr},0)", PCT2), ("H", f"=SUM(H5:H{lr-1})", MONEY),
                  ("I", f"=IFERROR(H{lr}/D{lr},0)", MONEY)]:
    cell(ws, f"{c}{lr}", f, fmt, font=bold)
ws.conditional_formatting.add(f"E5:E{lr-1}", DataBarRule(start_type="num", start_value=0, end_type="max", color=BLUE_HEX))
note(ws, f"A{lr + 2}", "12-month retention uses hires who started on or before 2024-12-31 (helper column retained_12m in Applications_Data).")
note(ws, f"A{lr + 3}", "Agency spend = 20% of first-year salary per hire; referral spend = $2,500 bonus per hire; other sources = subscriptions and events.")
ch = BarChart(); ch.type = "bar"; ch.title = "Application-to-hire rate by source"; ch.style = 10
ch.add_data(Reference(ws, min_col=5, min_row=4, max_row=lr - 2), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=lr - 2))
ch.x_axis.scaling.orientation = "maxMin"; ch.legend = None; ch.height, ch.width = 8, 15
ch.series[0].graphicalProperties.solidFill = BLUE_HEX
ws.add_chart(ch, f"A{lr + 5}")
ch2 = BarChart(); ch2.type = "bar"; ch2.title = "Cost per hire by source"; ch2.style = 10
ch2.add_data(Reference(ws, min_col=9, min_row=4, max_row=lr - 2), titles_from_data=True)
ch2.set_categories(Reference(ws, min_col=1, min_row=5, max_row=lr - 2))
ch2.x_axis.scaling.orientation = "maxMin"; ch2.legend = None; ch2.height, ch2.width = 8, 15
ch2.series[0].graphicalProperties.solidFill = "EB6834"
ws.add_chart(ch2, f"G{lr + 5}")

# ------------------------------------------------------------------ Offers
ws = wb.create_sheet("Offers", 3)
title(ws, "Why Are Offers Declined?", "Offer acceptance by department, salary gap and process length.")
header(ws, 4, 1, ["Department", "Offers decided", "Accepted", "Acceptance rate", "Avg offer"], [22, 13, 11, 13, 13])
for i, d in enumerate(DEPTS):
    r = 5 + i
    dc = f'{A_rng("department_name")},$A{r}'
    cell(ws, f"A{r}", d, font=bold)
    cell(ws, f"B{r}", f'=COUNTIFS({dc},{A_rng("offer_status")},"Accepted")+COUNTIFS({dc},{A_rng("offer_status")},"Declined")', INT)
    cell(ws, f"C{r}", f'=COUNTIFS({dc},{A_rng("offer_status")},"Accepted")', INT)
    cell(ws, f"D{r}", f"=IFERROR(C{r}/B{r},0)", PCT)
    cell(ws, f"E{r}", f'=IFERROR(AVERAGEIFS({A_rng("offered_salary")},{dc}),0)', MONEY)
ld = 5 + len(DEPTS) - 1
ws.conditional_formatting.add(f"D5:D{ld}", ColorScaleRule(start_type="min", start_color="EB6834",
                                                          mid_type="percentile", mid_value=50, mid_color="FFFFFF",
                                                          end_type="max", end_color="1BAF7A"))

reasons = sorted(apps.decline_reason.dropna().unique())
header(ws, 4, 7, ["Decline reason", "Offers declined", "Share"], [34, 13, 10])
for i, rs in enumerate(reasons):
    r = 5 + i
    cell(ws, f"G{r}", rs, font=bold)
    cell(ws, f"H{r}", f'=COUNTIFS({A_rng("decline_reason")},G{r})', INT)
    cell(ws, f"I{r}", f"=IFERROR(H{r}/SUM($H$5:$H${4 + len(reasons)}),0)", PCT)

r0b = ld + 3
header(ws, r0b, 1, ["Offer vs expected salary", "Lower bound", "Upper bound", "Offers", "Acceptance rate"])
bands = [(">10% below expected", -1, -0.10), ("5-10% below", -0.10, -0.05), ("0-5% below", -0.05, 0),
         ("At or above expected", 0, 1)]
for i, (lab, lo, hi) in enumerate(bands):
    r = r0b + 1 + i
    cell(ws, f"A{r}", lab, font=bold); cell(ws, f"B{r}", lo, PCT, fill=input_fill); cell(ws, f"C{r}", hi, PCT, fill=input_fill)
    rng = A_rng("offer_vs_expected_pct")
    base = f'{rng},">="&B{r},{rng},"<"&C{r}'
    cell(ws, f"D{r}", f'=COUNTIFS({base},{A_rng("offer_status")},"Accepted")+COUNTIFS({base},{A_rng("offer_status")},"Declined")', INT)
    cell(ws, f"E{r}", f'=IFERROR(COUNTIFS({base},{A_rng("offer_status")},"Accepted")/D{r},0)', PCT)
r0c = r0b + 7
header(ws, r0c, 1, ["Days application -> decision", "From (days)", "To (days)", "Offers", "Acceptance rate"])
for i, (lab, lo, hi) in enumerate([("Under 30 days", 0, 30), ("30-39 days", 30, 40), ("40-49 days", 40, 50), ("50+ days", 50, 999)]):
    r = r0c + 1 + i
    cell(ws, f"A{r}", lab, font=bold); cell(ws, f"B{r}", lo, INT, fill=input_fill); cell(ws, f"C{r}", hi, INT, fill=input_fill)
    rng = A_rng("days_in_process")
    base = f'{rng},">="&B{r},{rng},"<"&C{r}'
    cell(ws, f"D{r}", f'=COUNTIFS({base},{A_rng("offer_status")},"Accepted")+COUNTIFS({base},{A_rng("offer_status")},"Declined")', INT)
    cell(ws, f"E{r}", f'=IFERROR(COUNTIFS({base},{A_rng("offer_status")},"Accepted")/D{r},0)', PCT)
note(ws, f"A{r0c + 6}", "Yellow cells are band limits - change them to re-cut the analysis.")
for anchor, rr, ttl in ((f"G{r0b}", r0b, "Acceptance by salary vs expectation"),
                        (f"G{r0c + 8}", r0c, "Acceptance by process length")):
    ch = BarChart(); ch.type = "col"; ch.title = ttl; ch.style = 10
    ch.add_data(Reference(ws, min_col=5, min_row=rr, max_row=rr + 4), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=rr + 1, max_row=rr + 4))
    ch.legend = None; ch.height, ch.width = 7, 13; ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 1
    ch.y_axis.numFmt = "0%"; ch.series[0].graphicalProperties.solidFill = BLUE_HEX
    ws.add_chart(ch, anchor)

# ------------------------------------------------------------------ Recruiters
ws = wb.create_sheet("Recruiters", 4)
title(ws, "Recruiter Scorecard", "Workload, speed and results per recruiter.")
header(ws, 4, 1, ["Recruiter", "Requisitions", "Applications", "Avg resume-screen days", "Avg phone-screen days",
                  "Hires", "Avg time to hire", "Withdrawal rate", "Fill rate (closed reqs)"],
       [20, 12, 13, 14, 14, 9, 12, 12, 13])
for i, rc in enumerate(RECS):
    r = 5 + i
    cc = f'{A_rng("recruiter_name")},$A{r}'
    rq = f'{R_rng("recruiter_name")},$A{r}'
    cell(ws, f"A{r}", rc, font=bold)
    cell(ws, f"B{r}", f"=COUNTIFS({rq})", INT)
    cell(ws, f"C{r}", f"=COUNTIFS({cc})", INT)
    cell(ws, f"D{r}", f"=IFERROR(AVERAGEIFS({A_rng('days_resume_screen')},{cc}),0)", DEC)
    cell(ws, f"E{r}", f"=IFERROR(AVERAGEIFS({A_rng('days_phone_screen')},{cc}),0)", DEC)
    cell(ws, f"F{r}", f"=SUMIFS({A_rng('is_hired')},{cc})", INT)
    cell(ws, f"G{r}", f"=IFERROR(AVERAGEIFS({A_rng('time_to_hire_days')},{cc},{A_rng('is_hired')},1),0)", DEC)
    cell(ws, f"H{r}", f'=IFERROR(COUNTIFS({cc},{A_rng("outcome")},"Withdrew")/C{r},0)', PCT)
    cell(ws, f"I{r}", f'=IFERROR(COUNTIFS({rq},{R_rng("status")},"Filled")/(COUNTIFS({rq})-COUNTIFS({rq},{R_rng("status")},"Open")),0)', PCT)
lr = 5 + len(RECS) - 1
ws.conditional_formatting.add(f"G5:G{lr}", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="EB6834"))
ws.conditional_formatting.add(f"I5:I{lr}", ColorScaleRule(start_type="min", start_color="EB6834", end_type="max", end_color="FFFFFF"))
note(ws, f"A{lr + 2}", "Orange = needs attention (slow time to hire, low fill rate). Requisition counts cover Jan 2024 - Dec 2025.")

# ------------------------------------------------------------------ Monthly_Trend
ws = wb.create_sheet("Monthly_Trend", 5)
title(ws, "Monthly Trend", "Applications received and hires made each month.")
header(ws, 4, 1, ["Month", "Applications", "Hires", "Avg time to hire"], [12, 13, 9, 14])
months = sorted(apps.applied_month.unique())
for i, m in enumerate(months):
    r = 5 + i
    cell(ws, f"A{r}", m, font=bold)
    cell(ws, f"B{r}", f"=COUNTIFS({A_rng('applied_month')},$A{r})", INT)
    cell(ws, f"C{r}", f"=COUNTIFS({A_rng('outcome_month')},$A{r},{A_rng('is_hired')},1)", INT)
    cell(ws, f"D{r}", f'=IFERROR(AVERAGEIFS({A_rng("time_to_hire_days")},{A_rng("outcome_month")},$A{r},{A_rng("is_hired")},1),"")', DEC)
lm = 4 + len(months)
for col, ttl, anchor, color in ((2, "Applications per month", "F4", BLUE_HEX), (3, "Hires per month", "F20", "1BAF7A")):
    ch = LineChart(); ch.title = ttl; ch.style = 12; ch.height, ch.width = 7.5, 18
    ch.add_data(Reference(ws, min_col=col, min_row=4, max_row=lm), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=lm)); ch.legend = None
    ch.series[0].graphicalProperties.line.solidFill = color; ch.series[0].graphicalProperties.line.width = 22000
    ws.add_chart(ch, anchor)

# ------------------------------------------------------------------ Excel cleaning demo (raw rows)
ws = wb.create_sheet("Cleaning_Demo")
title(ws, "Excel Cleaning Demo", "A sample of RAW application rows with the formulas used to clean them in Excel.")
raw = pd.read_csv(BASE / "data" / "raw" / "applications.csv")
messy = raw[raw.applied_date.str.contains("/", na=False) | raw.source.isna()
            | ~raw.source.isin(SOURCES[:-1])].head(120)
dups = raw[raw.duplicated(keep=False)].head(30)
demo = pd.concat([raw.sample(150, random_state=1), messy, dups]).sample(frac=1, random_state=2)
demo = demo[["application_id", "source", "applied_date"]].reset_index(drop=True)
header(ws, 4, 1, ["application_id (raw)", "source (raw)", "applied_date (raw text)",
                  "Is duplicate?", "Source - trimmed", "Source - clean", "Applied date - clean"],
       [20, 22, 20, 12, 22, 22, 16])
map_start = 6
for i, row in enumerate(demo.itertuples(index=False)):
    r = 5 + i
    ws[f"A{r}"] = row.application_id; ws[f"B{r}"] = row.source if isinstance(row.source, str) else None
    ws[f"C{r}"] = str(row.applied_date)
    ws[f"D{r}"] = f'=IF(COUNTIFS($A$5:A{r},A{r},$C$5:C{r},C{r})>1,"Duplicate","")'
    ws[f"E{r}"] = f'=PROPER(TRIM(B{r}))'
    ws[f"F{r}"] = f'=IF(B{r}="","Unknown",IFERROR(INDEX($K${map_start}:$K${map_start + 20},MATCH(LOWER(TRIM(B{r})),$J${map_start}:$J${map_start + 20},0)),TRIM(B{r})))'
    ws[f"G{r}"] = (f'=IF(MID(C{r},3,1)="/",DATE(VALUE(RIGHT(C{r},4)),VALUE(LEFT(C{r},2)),VALUE(MID(C{r},4,2))),'
                   f'DATE(VALUE(LEFT(C{r},4)),VALUE(MID(C{r},6,2)),VALUE(RIGHT(C{r},2))))')
    ws[f"G{r}"].number_format = "yyyy-mm-dd"
    for c in "ABCDEFG": ws[f"{c}{r}"].font = norm
header(ws, map_start - 1, 10, ["Raw text (lower-case)", "Clean source"], [22, 20])
mapping = [("linkedin", "LinkedIn"), ("linked in", "LinkedIn"), ("job board", "Job Boards"), ("job boards", "Job Boards"),
           ("indeed/job boards", "Job Boards"), ("referral", "Employee Referral"), ("employee referral", "Employee Referral"),
           ("careers site", "Company Website"), ("company website", "Company Website"), ("recruitment agency", "Recruitment Agency"),
           ("campus recruiting", "Campus Recruiting"), ("social media", "Social Media")]
for i, (k, v) in enumerate(mapping):
    cell(ws, f"J{map_start + i}", k, fill=input_fill); cell(ws, f"K{map_start + i}", v)
note(ws, f"J{map_start + len(mapping) + 1}", "Mapping table used by the 'Source - clean' formula (INDEX/MATCH).")
note(ws, "I1", "Steps used on the full file: Remove Duplicates, TRIM/PROPER, a mapping table for")
note(ws, "I2", "inconsistent source names, and DATE(...) to rebuild US-style text dates. See Cleaning_Log.")

# ------------------------------------------------------------------ Cleaning_Log
ws = wb.create_sheet("Cleaning_Log")
title(ws, "Data Cleaning Log", "Every fix applied to the raw data before analysis (from scripts/02_clean_data.py).")
log = pd.read_csv(BASE / "data" / "clean" / "cleaning_log.csv").query("step != 'Loaded'")
header(ws, 4, 1, ["Table", "Step", "Rows affected", "Detail"], [16, 28, 14, 70])
for i, row in enumerate(log.itertuples(index=False)):
    r = 5 + i
    for j, v in enumerate(row):
        c = ws.cell(row=r, column=j + 1, value=v); c.font = norm; c.border = box
    ws.cell(row=r, column=3).number_format = INT

# ------------------------------------------------------------------ Dashboard (first sheet)
ws = wb.create_sheet("Dashboard", 0)
title(ws, "Recruitment Funnel Dashboard", "Jan 2024 - Dec 2025  |  all figures are live formulas over the data sheets")
for c, w in zip("ABCDEFGHIJK", [2, 20, 16, 2, 20, 16, 2, 20, 16, 2, 20]):
    ws.column_dimensions[c].width = w
hires_tot = f"SUM({A_rng('is_hired')})"
offer_acc = (f'COUNTIFS({A_rng("offer_status")},"Accepted")/(COUNTIFS({A_rng("offer_status")},"Accepted")'
             f'+COUNTIFS({A_rng("offer_status")},"Declined"))')
tiles = [  # (label cell, value cell, label, formula, number format)
    ("B4", "C4", "Requisitions", f"=COUNTA({R_rng('requisition_id')})", INT),
    ("E4", "F4", "Applications", f"=COUNTA({A_rng('application_id')})", INT),
    ("H4", "I4", "Hires", f"={hires_tot}", INT),
    ("B6", "C6", "Applications per hire", "=IFERROR(F4/I4,0)", INT),
    ("E6", "F6", "Avg time to hire (days)", f"=AVERAGEIFS({A_rng('time_to_hire_days')},{A_rng('is_hired')},1)", DEC),
    ("H6", "I6", "Avg time to fill (days)", f"=AVERAGE({R_rng('time_to_fill_days')})", DEC),
    ("B8", "C8", "Offer acceptance", f"={offer_acc}", PCT),
    ("E8", "F8", "Fill rate (closed reqs)",
     f'=COUNTIFS({R_rng("status")},"Filled")/(COUNTA({R_rng("requisition_id")})-COUNTIFS({R_rng("status")},"Open"))', PCT),
    ("H8", "I8", "Cost per hire", f"=SUM({S_rng('amount')})/I4", MONEY),
]
for lab_ref, val_ref, lab, f, fmt in tiles:
    cell(ws, lab_ref, lab, font=bold, fill=kpi_fill)
    c = cell(ws, val_ref, f, fmt, font=Font(name=F, bold=True, size=14, color=NAVY), fill=kpi_fill)
    c.alignment = Alignment(horizontal="right")
for r in (4, 6, 8): ws.row_dimensions[r].height = 24
ws["B10"] = "Key findings"; ws["B10"].font = Font(name=F, bold=True, size=12, color=NAVY)
findings = [
    "1. Onsite Interview is the bottleneck - the slowest stage, worst in Engineering and Data & Analytics (see Stage_Speed).",
    "2. Slow processes lose candidates: withdrawals rise sharply the longer candidates wait, and late offers are declined more (see Offers).",
    "3. About a third of offers are declined - mostly because pay is below the candidate's expectation.",
    "4. Employee referrals give the best conversion and retention at low cost; agencies cost ~8x more per hire (see Source_Analysis).",
    "5. Job boards and social media bring volume but very few hires.",
]
for i, t in enumerate(findings):
    ws[f"B{11 + i}"] = t; ws[f"B{11 + i}"].font = norm
note(ws, "B17", "Sheets: Funnel (interactive filters) | Stage_Speed | Source_Analysis | Offers | Recruiters | Monthly_Trend | data sheets | Cleaning_Demo | Cleaning_Log")

# funnel chart on dashboard (reads the Funnel sheet)
fs = wb["Funnel"]
ch = BarChart(); ch.type = "bar"; ch.title = "Hiring funnel (Funnel sheet filters apply)"; ch.style = 10
ch.add_data(Reference(fs, min_col=2, min_row=7, max_row=13), titles_from_data=True)
ch.set_categories(Reference(fs, min_col=1, min_row=8, max_row=13))
ch.x_axis.scaling.orientation = "maxMin"; ch.legend = None; ch.height, ch.width = 8, 14
ch.series[0].graphicalProperties.solidFill = BLUE_HEX
ws.add_chart(ch, "B19")
ss = wb["Stage_Speed"]
ch = BarChart(); ch.type = "col"; ch.title = "Average days per stage"; ch.style = 10
lastrow = 5 + len(DEPTS)
ch.add_data(Reference(ss, min_col=2, max_col=6, min_row=lastrow), from_rows=True, titles_from_data=False)
ch.set_categories(Reference(ss, min_col=2, max_col=6, min_row=4)); ch.legend = None; ch.height, ch.width = 8, 14
ch.series[0].graphicalProperties.solidFill = "EB6834"
ws.add_chart(ch, "F19")

# ------------------------------------------------------------------ README
ws = wb.create_sheet("README", 0)
title(ws, "Recruitment Funnel & Hiring Process Optimization - Excel workbook",
      "Stage 1 of the project (Excel -> SQL -> Python -> Power BI): first look at the data, cleaning and summaries.")
ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 100
rows = [
    ("Sheet", "What it shows"),
    ("Dashboard", "Headline KPIs, key findings and two summary charts."),
    ("Funnel", "Interactive funnel: choose a department and source in the yellow cells (dropdowns)."),
    ("Stage_Speed", "Average days in each stage by department - finds the bottleneck (colour scale)."),
    ("Source_Analysis", "Volume, conversion, spend, cost per hire, retention and performance by source."),
    ("Offers", "Acceptance by department, decline reasons, salary-gap bands and process-length bands (editable band limits)."),
    ("Recruiters", "Recruiter scorecard: workload, speed, hires, withdrawal and fill rates."),
    ("Monthly_Trend", "Applications and hires per month with line charts."),
    ("Applications_Data", "Clean data, one row per application (Excel Table 'Applications'). Helper column retained_12m is a formula."),
    ("Requisitions_Data", "One row per job opening (Excel Table 'Requisitions')."),
    ("Source_Spend", "Monthly recruiting spend by source (Excel Table 'SourceSpend')."),
    ("Cleaning_Demo", "Raw rows cleaned with Excel formulas: duplicate flag, TRIM/PROPER, mapping table, date rebuild."),
    ("Cleaning_Log", "Every cleaning step applied to the full raw data and how many rows it affected."),
    ("", ""),
    ("Formula notes", "Summaries use COUNTIFS, SUMIFS, AVERAGEIFS, INDEX/MATCH and IFERROR so they work in any Excel version."),
    ("Make a PivotTable", "Click any cell in Applications_Data > Insert > PivotTable. Try Rows = source, Values = Sum of is_hired."),
    ("Colour key", "Yellow cells = inputs you can change. Navy headers = calculated tables."),
]
for i, (a, b) in enumerate(rows):
    r = 4 + i
    ws[f"A{r}"], ws[f"B{r}"] = a, b
    ws[f"A{r}"].font = h_font if i == 0 else bold
    ws[f"B{r}"].font = h_font if i == 0 else norm
    if i == 0:
        ws[f"A{r}"].fill = ws[f"B{r}"].fill = h_fill

for w in wb.worksheets:
    w.sheet_properties.tabColor = {"README": "52514E", "Dashboard": NAVY}.get(w.title, None)
wb.active = 1  # open on the Dashboard
path = OUT / "recruitment_funnel_analysis.xlsx"
wb.save(path)
print("saved", path)
